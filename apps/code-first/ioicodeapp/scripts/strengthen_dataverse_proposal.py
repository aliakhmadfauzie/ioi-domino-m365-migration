import json
import os
import sys
from pathlib import Path
from urllib.parse import unquote

sys.path.insert(0, str(Path.cwd() / ".codex-docx-deps"))

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\aliak\OneDrive\Documents\IOI HCL Migration Projects")
INPUT = Path(r"C:\Users\aliak\Downloads\IOI_Dataverse_vs_SharePoint_Proposal.docx")
OUTPUT = Path(r"C:\Users\aliak\Downloads\IOI_Dataverse_vs_SharePoint_Proposal_Strengthened.docx")
ASSET_DIR = ROOT / "docs" / "proposal-assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)


def decode_site(value: str) -> str:
    return unquote(unquote(value))


def read_sharepoint_inventory():
    config = json.loads((ROOT / "powerapps-mirror/apps/ioi-platform/power.config.json").read_text())
    seen = {}
    for ref_id, ref in config.get("connectionReferences", {}).items():
        if ref.get("id") != "/providers/Microsoft.PowerApps/apis/shared_sharepointonline":
            continue
        for site_key, site_data in ref.get("dataSets", {}).items():
            site = decode_site(site_key)
            for list_name, details in site_data.get("dataSources", {}).items():
                table_id = details.get("tableName", "")
                key = (site, list_name, table_id)
                seen.setdefault(key, {"site": site, "list": list_name, "table_id": table_id, "ref": ref_id})
    rows = sorted(seen.values(), key=lambda x: (x["site"].lower(), x["list"].lower()))
    groups = {}
    for row in rows:
        groups.setdefault(row["site"], []).append(row)
    return rows, groups


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, size=9, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, header_fill="E8EEF5"):
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for row_i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(8.5)
            if row_i == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(31, 77, 120)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    try:
        p.style = f"Heading {level}"
    except KeyError:
        pass
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16 if level == 1 else 13 if level == 2 else 12)
    run.font.color.rgb = RGBColor(31, 77, 120)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    if "Normal" in [style.name for style in doc.styles]:
        p.style = "Normal"
    p.paragraph_format.space_after = Pt(6)
    p.add_run(text)
    return p


def add_bullet(doc, text):
    style_names = {style.name for style in doc.styles}
    p = doc.add_paragraph(style="List Bullet" if "List Bullet" in style_names else None)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)
    return p


def draw_box(ax, xy, wh, text, fc, ec="#1F4D78", size=9, weight="normal"):
    x, y = xy
    w, h = wh
    rect = plt.Rectangle((x, y), w, h, facecolor=fc, edgecolor=ec, linewidth=1.2)
    ax.add_patch(rect)
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=size, weight=weight, wrap=True)


def arrow(ax, start, end):
    ax.annotate("", xy=end, xytext=start, arrowprops=dict(arrowstyle="->", lw=1.2, color="#34495E"))


def make_current_architecture_diagram(groups):
    path = ASSET_DIR / "current-sharepoint-architecture.png"
    fig, ax = plt.subplots(figsize=(10, 6), dpi=180)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis("off")

    draw_box(ax, (0.3, 4.6), (2.0, 0.8), "Power Apps\nCanvas + Code App", "#DCEBFA", weight="bold")
    draw_box(ax, (0.3, 3.1), (2.0, 0.8), "Power Automate\nFlows + email", "#EEF4FB")
    draw_box(ax, (0.3, 1.6), (2.0, 0.8), "Power BI / Reports\ncross-site joins", "#EEF4FB")
    draw_box(ax, (3.2, 4.85), (2.1, 0.55), "SharePoint connector\nmultiple connection refs", "#FFF2CC", "#B68A00", weight="bold")
    draw_box(ax, (3.2, 3.9), (2.1, 0.55), "Full estate\n330-440 physical lists", "#FFF8DB", "#B68A00")
    draw_box(ax, (3.2, 2.9), (2.25, 0.68), "22 departments\napprox. 418\nlist-equivalents", "#FFF8DB", "#B68A00", size=8.1)

    sites = [
        ("IT benchmark\n19 business lists", 6.2, 4.75, "#FADBD8"),
        ("HR / QA / SHE\nsimilar patterns", 8.0, 4.75, "#FADBD8"),
        ("22 department\nsite/list estates", 6.2, 3.75, "#FDEDEC"),
        ("Approval + audit\nlists per dept", 8.0, 3.75, "#FDEDEC"),
        ("Extension/detail\nlists per form family", 6.2, 2.75, "#FDEDEC"),
        ("Config/routing\nreference lists", 8.0, 2.75, "#FDEDEC"),
        ("Admin/security\nshared platform lists", 7.1, 1.75, "#FDEDEC"),
    ]
    for label, x, y, fill in sites:
        draw_box(ax, (x, y), (1.45, 0.62), label, fill, "#922B21", size=8.2)

    for y in (5.0, 3.5, 2.0):
        arrow(ax, (2.3, y), (3.2, 5.1 if y == 5.0 else 4.17 if y == 3.5 else 3.22))
    arrow(ax, (5.3, 5.1), (6.2, 5.05))
    arrow(ax, (5.3, 4.17), (6.2, 4.06))
    arrow(ax, (5.45, 3.24), (6.2, 3.06))

    ax.text(
        0.4,
        0.45,
        "Full-estate issue: requests, extension details, approvals, documents, routing, audit, and security are distributed across hundreds of SharePoint lists.",
        fontsize=9,
        color="#333333",
    )
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def make_dataverse_diagram():
    path = ASSET_DIR / "proposed-dataverse-architecture.png"
    fig, ax = plt.subplots(figsize=(10, 6), dpi=180)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis("off")

    draw_box(ax, (0.4, 4.6), (2.0, 0.75), "Canvas Apps\nCode Apps\nPower Pages (future)", "#DCEBFA", weight="bold")
    draw_box(ax, (0.4, 3.35), (2.0, 0.75), "Power Automate\nApprovals + SLA", "#EEF4FB")
    draw_box(ax, (0.4, 2.1), (2.0, 0.75), "Power BI\nmodel-driven reporting", "#EEF4FB")
    draw_box(ax, (3.2, 2.45), (2.0, 2.25), "Dataverse Environment\nDEV / TEST / PROD\nSolutions + ALM\nSecurity roles + Teams", "#D5F5E3", "#196F3D", weight="bold")

    tables = [
        ("cr_request\none row per form", 6.1, 4.9),
        ("cr_approvalrecord\nstage rows", 8.0, 4.9),
        ("cr_requestextension\nform-specific details", 6.1, 3.85),
        ("cr_attachment\nmetadata + link", 8.0, 3.85),
        ("cr_routingrule\ncompany/site/team", 6.1, 2.8),
        ("cr_slaevent\nclock + escalation", 8.0, 2.8),
        ("cr_lookupvalue\ncompany/dept/system", 6.1, 1.75),
        ("Dataverse audit\nfield-level history", 8.0, 1.75),
    ]
    for label, x, y in tables:
        draw_box(ax, (x, y), (1.55, 0.58), label, "#EAF8EA", "#196F3D", size=8.1)

    for y in (4.98, 3.73, 2.48):
        arrow(ax, (2.4, y), (3.2, 3.58))
    arrow(ax, (5.2, 3.6), (6.1, 5.1))
    arrow(ax, (5.2, 3.6), (6.1, 4.08))
    arrow(ax, (5.2, 3.6), (6.1, 3.05))
    arrow(ax, (7.65, 5.1), (8.0, 5.1))
    arrow(ax, (7.65, 4.08), (8.0, 4.08))
    arrow(ax, (7.65, 3.05), (8.0, 3.05))

    ax.text(
        0.45,
        0.55,
        "Target principle: SharePoint becomes a document store where useful; Dataverse becomes the transactional system of record.",
        fontsize=9,
        color="#333333",
    )
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def add_inventory_table(doc, groups):
    table = doc.add_table(rows=1, cols=4)
    headers = ["SharePoint site", "Configured lists/bindings", "Representative lists", "Architecture concern"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8.5, color="1F4D78")
    concerns = {
        "ioi-portal-it": "Highest coupling: parent request, extension tables, routing, approvals, policies, and reference data.",
        "ioi-portal-hr": "HR records, health/appraisal, audit, attachments, and routing split across multiple lists.",
        "ioi-portal-qa": "Change/action records require cross-form traceability and retention.",
        "IOI-App_Admin_Sites": "Platform security/admin data lives outside business sites.",
        "MainDB_IT": "Routing matrix is isolated from IT transaction lists.",
    }
    for site, rows in sorted(groups.items(), key=lambda x: x[0].lower()):
        cells = table.add_row().cells
        names = [r["list"] for r in rows]
        site_short = site.replace("https://ioioi.sharepoint.com/sites/", "")
        concern = next((v for k, v in concerns.items() if k in site), "Separate site/list boundary adds permission and reporting overhead.")
        set_cell_text(cells[0], site_short, size=8.2)
        set_cell_text(cells[1], len(rows), size=8.2)
        set_cell_text(cells[2], ", ".join(names[:5]) + (" ..." if len(names) > 5 else ""), size=8.0)
        set_cell_text(cells[3], concern, size=8.0)
    style_table(table)
    return table


def add_full_estate_architecture_table(doc):
    rows = [
        [
            "Department transaction lists",
            "Parent request lists plus form-family extension/detail lists",
            "Approximately 418 departmental list-equivalents",
            "This is the main risk driver: request data is fragmented across parent/detail lists instead of normalized Dataverse relationships.",
        ],
        [
            "Physical SharePoint list estate",
            "Department list-equivalents after consolidation/shared lists are accounted for",
            "330-440 physical SharePoint lists",
            "Use this as the full-estate support and governance range for permissions, schema changes, testing, and reporting.",
        ],
        [
            "Completed benchmark",
            "IT department pattern",
            "19 practical IT business lists, plus 1 generated lookup alias",
            "IT is the benchmark because it includes parent request, extension, approval, workflow, routing, reference, policy, and role lists.",
        ],
        [
            "Shared platform lists",
            "Admin, roles, route permissions, audit, app settings, user profiles, global configuration",
            "Included in the 330-440 physical-list range",
            "These lists create cross-cutting dependencies across departments and should not be treated as isolated department data.",
        ],
        [
            "Testing/build evidence",
            "Live current power.config.json connection map",
            "51 bindings across 9 sites in the tested build snapshot",
            "This proves the architecture pattern and testing findings; it is not used as the full-estate count.",
        ],
    ]
    table = doc.add_table(rows=1, cols=4)
    headers = ["Architecture layer", "What it represents", "Count used in proposal", "Why it matters"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8.2, color="1F4D78")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text, size=7.8)
    style_table(table)
    return table


def add_completion_projection_table(doc):
    rows = [
        [
            "Tested build snapshot",
            "Live power.config.json from the completed SharePoint build",
            "51 SharePoint bindings across 9 sites",
            "This is the tested build snapshot used as evidence of the architecture pattern; it is not the full-estate count.",
        ],
        [
            "Completed department example: IT",
            "ioi-portal-it site",
            "20 bindings, including 1 generated lookup alias; 19 practical business lists",
            "Use IT as the completed-department benchmark because it has parent, extension, approval, workflow, config, routing, reference, policy, and role lists.",
        ],
        [
            "Simple extrapolation to 22 departments",
            "19 practical departmental list-equivalents x 22 departments",
            "Approximately 418 departmental list-equivalents",
            "This uses IT as the completed department benchmark for full-rollout sizing.",
        ],
        [
            "Physical list range after consolidation",
            "Department list-equivalents plus shared/admin/global lists",
            "330-440 physical SharePoint lists",
            "This is the supportable full-estate range after shared lists and consolidation are accounted for.",
        ],
        [
            "Full-estate physical list range",
            "Completed estate sizing using department blueprint pattern",
            "330-440 SharePoint lists",
            "Use this as the full-estate SharePoint list range for impact, governance, and support sizing.",
        ],
    ]
    table = doc.add_table(rows=1, cols=4)
    headers = ["Rollout layer", "Basis", "Projected count", "How to use it"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8.2, color="1F4D78")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text, size=7.8)
    style_table(table)
    return table


def add_failure_table(doc):
    rows = [
        [
            "Schema/write failure",
            'ITSSR extension submit failed when the generated SharePoint payload sent a field shape that the list did not accept. A prior failure was "The passed-in field title could not be found". Another was CurrentStatus#Id, which was readable/internal but not writable.',
            "This is not just slow performance. It is a reliability failure caused by SharePoint connector metadata, generated TypeScript models, and actual list schema drifting apart.",
            "Dataverse tables expose stable typed columns and relationships through solutions. Choice/status fields are written as Dataverse values, not connector-specific SharePoint read/write shapes.",
        ],
        [
            "Relationship workaround",
            "ITSSR must create a parent row in maindb_it, then create a child row in maindb_it_ext_support using ParentID plus a title/case-number fallback.",
            "The app has to defend against lookup and title mismatches because SharePoint is acting as a relational database without relational guarantees.",
            "Use cr_request as the parent table and cr_requestextension / cr_itssr_detail as related child rows with enforced Dataverse lookup relationships.",
        ],
        [
            "Performance/data loading",
            "The ITSSR list screen needed a second hardening pass to add top/skip paging and to narrow child-list queries to only the loaded parent records.",
            "The original SharePoint pattern scales by manually stitching lists and limiting each query. Search/filter behavior can become page-local unless more server-side query logic is built.",
            "Dataverse supports indexed server-side filtering, relationship-aware retrieval, views, and Power BI-friendly reporting without cross-site list stitching.",
        ],
        [
            "Permissions at scale",
            "A direct user-to-list permission model across 50 configured lists and 54 users timed out after 30 minutes; the workable approach had to pivot to per-site SharePoint groups.",
            "This shows operational fragility. Permissions are spread over sites and lists instead of being managed as business roles around records and teams.",
            "Dataverse security roles, owner teams, access teams, business units, and column security map directly to requestor, approver, reader, IT PIC, and admin roles.",
        ],
    ]
    table = doc.add_table(rows=1, cols=4)
    headers = ["Observed issue", "Where/how it failed", "Why it matters", "Dataverse design response"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8.2, color="1F4D78")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text, size=7.6)
    style_table(table)
    return table


def add_dataverse_table_architecture(doc):
    rows = [
        ["cr_application", "One row per migrated Domino application/form", "Department, form code, route, retention class, owner team", "1:N cr_request, cr_routingrule"],
        ["cr_request", "Main transaction record replacing maindb_* lists", "Request no, title, status, workflow stage, requestor, company, department, site, submitted date", "1:N approval, attachment, SLA, audit, extension"],
        ["cr_requestextension / form-specific tables", "Normalized form details for ITSSR, EAF, IAL, HR Interview, QA CC, etc.", "Service type, hardware, application, SAP client, candidate score, line-item fields", "N:1 cr_request; use one table per complex domain where fields differ materially"],
        ["cr_approvalrecord", "Every approval step is a row, not a set of columns", "Stage no, role, approver, delegated approver, status, decision date, comments, SLA due date", "N:1 cr_request; N:1 cr_routingrule"],
        ["cr_routingrule", "Configurable routing by department/company/site/form", "Form code, company, site, role, approver team, escalation team, active dates", "Referenced by approvals and flows"],
        ["cr_slaevent", "SLA clock, reminder, escalation, and breach history", "Start/end, due date, paused time, escalated to, breach reason", "N:1 cr_request and approval stage"],
        ["cr_attachment", "Attachment metadata; file can remain in SharePoint/Dataverse file column by policy", "Document type, file link, uploaded by/date, parent request, version", "N:1 cr_request"],
        ["cr_auditentry", "Structured business audit beyond platform audit", "Action, actor, old value, new value, source app, timestamp", "N:1 cr_request"],
        ["cr_lookupvalue", "Shared reference data replacing scattered config lists", "Company, department, system, bank, SAP client, server, policy type", "Referenced by request/detail/routing tables"],
        ["cr_securityassignment", "Domino reader/editor equivalent", "User/team, access role, scope, start/end date", "N:1 request/application; maps to Dataverse teams/roles"],
    ]
    table = doc.add_table(rows=1, cols=4)
    headers = ["Dataverse table", "Purpose", "Core columns", "Relationships"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8.2, color="1F4D78")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text, size=7.7)
    style_table(table)
    return table


def add_migration_mapping(doc):
    rows = [
        ["maindb_it, maindb_hr, maindb_qa, maindb_str, maindb_she, etc.", "cr_request", "One common parent table with department/form discriminator. Keeps reporting consistent across departments."],
        ["maindb_it_ext_support, maindb_it_ext_access, maindb_it_ext_sap, QA sub/action lists, STR line items", "cr_requestextension or typed child tables", "Use typed child tables where each form family has unique fields; avoid one mega-table for unrelated detail fields."],
        ["approvalrecords, approvalhistory_*", "cr_approvalrecord", "Approval stage as rows. Adding a stage becomes configuration, not a schema change."],
        ["workflowauditlog, auditlogs, auditlog_hr", "cr_auditentry + Dataverse audit", "Separate business audit from platform audit and make it queryable."],
        ["config_* lists, routingmatrix lists, it_m1_rs_* reference lists", "cr_lookupvalue, cr_routingrule, cr_slasetting", "Centralize configuration and route rules so Code Apps, Canvas Apps, and flows share one source of truth."],
        ["attachmentindex_* and document extension lists", "cr_attachment + SharePoint document library or Dataverse file column", "Store metadata relationally; keep large binary storage policy-driven."],
        ["userroles, routepermissions, user_roles_it", "Dataverse security roles/teams + cr_securityassignment", "Use platform security first, with explicit assignment table only where Domino reader/editor migration requires it."],
    ]
    table = doc.add_table(rows=1, cols=3)
    headers = ["Current SharePoint source", "Dataverse target", "Migration rationale"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=8.2, color="1F4D78")
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text, size=7.8)
    style_table(table)
    return table


def configure_styles(doc):
    styles = doc.styles
    try:
        styles["Normal"].font.name = "Calibri"
        styles["Normal"].font.size = Pt(10.5)
    except KeyError:
        pass
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        try:
            styles[name].font.color.rgb = RGBColor(31, 77, 120)
        except KeyError:
            pass


def remove_closing_footer(doc):
    removed = []
    markers = (
        "Migration Programme Office  |  IOI Oleochemical Sdn Bhd  |  May 2026",
        "Document prepared in support of the Lotus Domino to Microsoft Power Platform Migration Programme.",
    )
    for p in list(doc.paragraphs):
        if p.text.strip() in markers:
            removed.append(p.text.strip())
            p._element.getparent().remove(p._element)
    return removed


def add_closing_footer(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    run = p.add_run("Migration Programme Office  |  IOI Oleochemical Sdn Bhd  |  May 2026")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("Document prepared in support of the Lotus Domino to Microsoft Power Platform Migration Programme.")
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(89, 89, 89)


def replace_paragraph_text(paragraph, text):
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def apply_fact_check_replacements(doc):
    replacements = {
        "Prepared by: Migration Programme Office  |  May 2026": (
            "Prepared by: Ali Akhmad Fauzie  |  May 2026"
        ),
        "SharePoint Lists are flat by design. They do not support true relational data. Lookup columns exist but are limited: no cascading deletes, no referential integrity, no multi-column foreign keys, and a hard lookup threshold of 12 lookup columns per list before performance degrades significantly.": (
            "SharePoint Lists are list-based rather than relational-database tables. They can use lookup columns, unique columns, and limited enforced relationship behavior, but those relationships are constrained by list-view and lookup thresholds, do not provide Dataverse-style relational modelling, and become difficult to govern across many applications and site collections."
        ),
        "SharePoint cannot do this. SharePoint security is at the list level, or at the item level only through unique permissions — which are explicitly banned by Microsoft for lists with more than 50,000 items due to performance, and are operationally unmanageable at the scale of 361 applications.": (
            "SharePoint can apply unique permissions at item level, but Microsoft documents a supported limit of 50,000 unique permission scopes per list or library and recommends keeping the general limit at 5,000 for performance. That makes item-level security a poor fit for reproducing Domino Reader/Editor semantics across 361 applications."
        ),
        "Limited formula column — no IF/AND/OR nesting": (
            "Calculated columns exist, but complex business rules and workflow-state logic become difficult to govern across lists"
        ),
        "No native retention policy — must rely on SharePoint retention labels": (
            "Retention is handled through Microsoft Purview labels/policies rather than transaction-table rules tied to a relational request model"
        ),
        "Dataverse tables have no item threshold. Indexed columns and server-side filtering scale to tens of millions of rows.": (
            "Dataverse is not constrained by the SharePoint 5,000-item list-view threshold. It scales through indexed server-side queries, purchased storage capacity, and Power Platform API/service-protection governance."
        ),
        "HIGH — Domino Reader/Editor fields cannot be faithfully migrated to SharePoint without unique item permissions, which degrade performance at >5,000 items": (
            "HIGH — Domino Reader/Editor fields can only be approximated in SharePoint through unique item permissions or application-layer filtering; Microsoft recommends staying under 5,000 unique scopes for performance"
        ),
        "MEDIUM — SharePoint retention labels apply at the library level, not the list item level. Fine-grained ISO 9001 retention enforcement is manual": (
            "MEDIUM — SharePoint retention is governed through Microsoft Purview labels/policies and site/list content controls, but it is not a normalized transaction-table retention model tied to request, approval, SLA, and child-record relationships"
        ),
        "Dataverse retention policies apply at the table level with field-level triggers. Audit log is queryable and exportable.": (
            "Dataverse long-term retention policies can be configured for tables and related child data, subject to supported table types and platform capacity/API rules. Dataverse auditing provides structured row and field history for governed records."
        ),
        "Dataverse for Teams is included in Microsoft 365 E1/E3/E5 licences at no additional cost for apps built with Power Apps for Teams.": (
            "Dataverse for Teams is included with many Microsoft 365/Office 365 subscriptions that include Teams and Power Platform capabilities, but it has environment and capacity limits such as 2 GB per team and up to 1 million rows. It is suitable for Teams-scoped apps, not for every enterprise production scenario."
        ),
        "Standalone Dataverse (for apps published outside Teams) requires Power Apps Per User or Per App licence.": (
        "Apps using standalone Dataverse or running outside the Dataverse for Teams use rights generally require appropriate Power Apps premium licensing, such as per app or per user licensing."
        ),
        "The IOI Oleochemical Lotus Domino to Microsoft Power Platform migration covers approximately 361 applications across 22 departments at two sites — Penang and Johor. We are mid-programme. The architecture decisions made now will either scale cleanly to full deployment or accumulate technical debt that will cost significantly more to resolve later.": (
            "The IOI Oleochemical Lotus Domino to Microsoft Power Platform migration covers approximately 361 applications across 22 departments at two sites — Penang and Johor. The SharePoint-based implementation work has progressed to completion for the SharePoint build, and testing has now exposed architecture-level issues that must be resolved before production scale-out. The decision now is whether to remediate those issues at the data-layer level with Dataverse or continue absorbing SharePoint-specific workarounds."
        ),
        "This document presents a data-driven case — drawn directly from the application forms already analysed in this programme — for using Microsoft Dataverse as the primary data layer, rather than SharePoint Lists, for the majority of applications in scope.": (
            "This document presents a testing-driven recommendation — drawn directly from the completed SharePoint build, full-estate sizing, and application forms already analysed in this programme — for using Microsoft Dataverse as the primary data layer rather than continuing to operate the SharePoint List architecture as the enterprise system of record."
        ),
        "This is not a theoretical recommendation. Every limitation cited below has already been encountered in real applications from this migration backlog.": (
            "This is not a theoretical recommendation. The issues cited below were found during implementation and testing: schema/write failures, lookup relationship fragility, list-screen performance hardening, permission-scale friction, and reporting/audit complexity."
        ),
        "What Happens If We Don't Change": (
            "What Happens If We Keep the Completed SharePoint Build Unchanged"
        ),
        "Continuing to build all applications on SharePoint Lists — without a formal decision to use Dataverse for complex forms — will produce the following foreseeable outcomes:": (
            "The completed SharePoint-based build can be made to run for selected forms, but testing shows that keeping it unchanged as the enterprise data architecture will produce the following outcomes:"
        ),
        "HIGH — IT ITSSR, QA Change Control, and HR recruitment logs will all exceed 5,000 items within 12-18 months of go-live": (
            "HIGH — IT ITSSR, QA Change Control, and HR recruitment logs are expected to exceed SharePoint list-view operating thresholds within 12-18 months of production usage unless queries, indexes, and data partitioning are tightly managed"
        ),
        "Re-work cost when SharePoint limits are hit post-launch": (
            "Re-work cost if testing findings are carried into production"
        ),
        "CERTAIN — Every application hitting these limits must be re-architected after users are already in production": (
            "CERTAIN — Every application that grows beyond the tested SharePoint workarounds will need remediation after users are already relying on it"
        ),
        "Architectural investment in Dataverse is made once, before production. No forced re-architecture mid-flight.": (
            "Dataverse remediation is made before production scale-out, avoiding repeated post-testing fixes and production re-architecture."
        ),
        "The case for Dataverse should be framed as more than 'SharePoint may be slow'. In this programme, SharePoint has already produced both hard failures and performance/maintainability issues. The failures came from schema drift, connector read/write differences, manual relationship handling, and permission operations spread across many lists.": (
            "The case for Dataverse should be framed as more than 'SharePoint may be slow'. The completed SharePoint build exposed both hard failures and performance/maintainability issues during testing. The failures came from schema drift, connector read/write differences, manual relationship handling, and permission operations spread across many lists."
        ),
    }

    table_replacements = {
        "Lookup column limited to 12 per list before performance drops": (
            "Lookup/view joins are constrained by SharePoint lookup/list-view thresholds"
        ),
        "Unique item permissions — banned at scale by Microsoft": (
            "Unique item permissions supported but capped/recommended below threshold"
        ),
        "Native Business Unit / Team ownership — scales to millions": (
            "Native Business Unit / Team ownership governed by Dataverse capacity/API limits"
        ),
        "No native YAML export — PowerShell only": (
            "No Dataverse-style solution package/YAML ALM model for list schemas"
        ),
        "YAML schema export, solution packaging, Git-deployable": (
            "Solution packaging and YAML source-control format available with current Power Platform tooling"
        ),
    }

    all_paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)

    for paragraph in all_paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            replace_paragraph_text(paragraph, replacements[text])

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text in table_replacements:
                    set_cell_text(cell, table_replacements[text], size=7.8)


def main():
    rows, groups = read_sharepoint_inventory()
    current_diagram = make_current_architecture_diagram(groups)
    dataverse_diagram = make_dataverse_diagram()

    doc = Document(INPUT)
    configure_styles(doc)
    remove_closing_footer(doc)
    apply_fact_check_replacements(doc)

    add_heading(doc, "POST-BUILD TESTING POSITION", 1)
    add_body(
        doc,
        "The recommendation to move complex applications to Dataverse is based on findings from the completed SharePoint-based build and subsequent testing. "
        "The issue is not that the team could not build the screens or connect the lists. The issue is that testing exposed data-layer weaknesses that will grow worse as the full estate moves into production usage.",
    )
    add_body(
        doc,
        "The SharePoint build proved the business process and user interface direction, but it also showed that enterprise-scale Domino replacement requires relational integrity, governed row-level security, stable schema deployment, structured audit, and scalable reporting. "
        "Those are Dataverse responsibilities, not SharePoint List workarounds.",
    )

    add_heading(doc, "APPENDIX A - SHAREPOINT ESTATE ARCHITECTURE", 1)
    add_body(
        doc,
        "The completed SharePoint architecture should be assessed as a full estate, not as a single-list or single-department pattern. "
        "Using the completed IT department as the benchmark, the rollout represents approximately 418 departmental list-equivalents and an estimated 330-440 physical SharePoint lists after shared/admin/global lists and consolidation are accounted for.",
    )
    doc.add_picture(str(current_diagram), width=Inches(6.4))
    add_caption(doc, "Figure A1. Full SharePoint estate architecture using the completed IT department as the benchmark.")
    add_full_estate_architecture_table(doc)
    add_caption(doc, "Table A1. Full-estate SharePoint architecture count used in this proposal.")

    add_heading(doc, "How the Lists Are Connected Today", 2)
    add_body(
        doc,
        "The completed SharePoint architecture uses lists as if they were relational tables. A typical IT form such as ITSSR is split across maindb_it for the parent request, "
        "maindb_it_ext_support for service-request details, document/attachment lists for files, workflow/audit lists for history, and separate config/routing lists for approver selection. "
        "The connection is enforced by application code using ParentID lookups, case-number/title matching, and OData filters rather than by database constraints.",
    )
    for text in [
        "Parent request lists hold shared workflow fields such as request number, status, requestor, submitted date, and workflow stage.",
        "Extension lists hold form-specific details such as ITSSR service type, hardware/application, assigned IT PIC, solution, satisfaction rating, and PanCen approval fields.",
        "Approval and audit lists record actions separately from the transaction record, which makes reporting and retention depend on cross-list joins.",
        "Config and routing lists hold companies, departments, approver matrices, routing rules, systems, banks, SAP clients, policies, and user roles.",
        "Each site collection has its own permission boundary, so security is managed across sites/lists rather than as one coherent business data model.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "Full Rollout SharePoint List Projection", 2)
    add_body(
        doc,
        "For full-estate sizing, the IT department is used as the completed department benchmark. IT shows 20 SharePoint bindings on the IT site, including one generated lookup alias that points back to maindb_it; therefore the practical IT business-list benchmark is 19 list-equivalents.",
    )
    add_body(
        doc,
        "Applying that completed-department pattern across 22 departments results in roughly 418 departmental list-equivalents. After shared/admin/global lists and consolidation are accounted for, use a total physical SharePoint estate of approximately 330-440 lists for governance, support, permissions, reporting, and data-layer risk sizing.",
    )
    add_completion_projection_table(doc)
    add_caption(doc, "Table A2. SharePoint full-rollout list-count projection using IT as the completed department benchmark.")

    add_heading(doc, "APPENDIX B - WHERE SHAREPOINT DID NOT WORK CLEANLY", 1)
    add_body(
        doc,
        "The case for Dataverse should be framed as more than 'SharePoint may be slow'. In this programme, SharePoint has already produced both hard failures and performance/maintainability issues. "
        "The failures came from schema drift, connector read/write differences, manual relationship handling, and permission operations spread across many lists.",
    )
    add_failure_table(doc)
    add_caption(doc, "Table B1. Evidence from the current implementation and repair history.")

    add_heading(doc, "Specific ITSSR Evidence", 2)
    for text in [
        "Submit reliability: ITSSR creates a parent record and then a child extension record. When generated payload fields did not match the live SharePoint schema, submit failed with SharePoint 400 errors instead of a recoverable business validation message.",
        "Status writes: CurrentStatus had to be written as a plain text/status field; attempts to write the connector's read-side CurrentStatus#Id shape failed because that field is internal/read-only.",
        "Parent-child lookup fragility: child records are found by ParentID/Id and title fallback. This is code-level relationship stitching, not referential integrity.",
        "List-screen performance: the app needed explicit page-size controls, top/skip paging, and child-list narrowing by parent records. This improved the screen, but it is a mitigation for the SharePoint-list shape, not a long-term data architecture.",
        "Permission operations: granting list access at scale had to pivot from direct per-user/per-list grants to per-site groups after timeout/friction. Dataverse handles this as roles and teams around the data model.",
    ]:
        add_bullet(doc, text)

    add_heading(doc, "APPENDIX C - PROPOSED DATAVERSE TABLE ARCHITECTURE", 1)
    add_body(
        doc,
        "The proposed Dataverse architecture keeps the application experience familiar while replacing scattered SharePoint lists with a normalized, governed data layer. "
        "Canvas Apps and Code Apps continue to provide the user interface; Power Automate continues to handle approvals, notifications, and integrations; Dataverse becomes the transactional system of record.",
    )
    doc.add_picture(str(dataverse_diagram), width=Inches(6.4))
    add_caption(doc, "Figure C1. Proposed Dataverse architecture for IOI Power Platform applications.")
    add_dataverse_table_architecture(doc)
    add_caption(doc, "Table C1. Proposed Dataverse table architecture.")

    add_heading(doc, "SharePoint-to-Dataverse Mapping", 2)
    add_migration_mapping(doc)
    add_caption(doc, "Table C2. Migration mapping from the completed SharePoint estate to Dataverse tables.")

    add_heading(doc, "Implementation Principle", 2)
    add_body(
        doc,
        "SharePoint should remain available where it is strongest: document libraries, simple low-volume lists, and user-facing collaboration content. "
        "Dataverse should own transactional records, approvals, relationships, audit, security, routing, SLA, and reporting. This separation avoids rebuilding Domino complexity inside SharePoint workarounds.",
    )

    add_closing_footer(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
